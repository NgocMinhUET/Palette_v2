# coding=utf-8
"""Insert intro sentences referencing each figure/table that lacks a body reference.

Strategy: for every caption flagged MISSING in the v12 document, insert a new
paragraph immediately BEFORE the caption (and before the figure placeholder
if present) that introduces the figure/table in the body text.
"""

import re
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

SRC = r"C:\Users\Lenovo\Downloads\Tieu_luan_cuoi_khoa_QoE_cross_layer_VCM_DoNgocMinh_v12.docx"
OUT = r"C:\Users\Lenovo\Downloads\Tieu_luan_cuoi_khoa_QoE_cross_layer_VCM_DoNgocMinh_v12.docx"  # overwrite

# Intro sentences keyed by "Kind Number"
INTROS = {
    "Hình 1": (
        "Để minh họa rõ trực giác này, Hình 1 dưới đây trình bày vai trò bổ "
        "sung lẫn nhau giữa application layer và network-side delivery trong "
        "bài toán đang xét: application layer quyết định phần thông tin nào "
        "cần được giữ lại, còn network-side delivery quyết định phần thông "
        "tin đó được phân phối ra sao trong điều kiện mạng cụ thể."
    ),
    "Hình 2": (
        "Bức tranh tổng quát về các nhánh nghiên cứu liên quan và vị trí học "
        "thuật của đề tài được khái quát ở Hình 2, trong đó network-aware VCM "
        "được đặt ở giao điểm của ba mảnh ghép chính: VCM nền tảng, "
        "task-oriented communication và cross-layer QoE optimization."
    ),
    "Bảng 2": (
        "Để hệ thống hóa vai trò của các công trình VCM nền tảng đối với bài "
        "toán đang xét, Bảng 2 dưới đây tổng hợp đóng góp, trọng tâm, mức độ "
        "network-awareness và ý nghĩa của từng công trình trong khung đề tài."
    ),
    "Hình 3": (
        "Để hình dung cách ba thành phần của khung TOCOM-TEM phối hợp với "
        "nhau, Hình 3 trình bày sơ đồ khái niệm của framework này, từ "
        "task-relevant feature extraction tại thiết bị, qua temporal entropy "
        "model, đến spatial-temporal fusion ở edge server."
    ),
    "Hình 4": (
        "Để mô tả trực quan cách kiến trúc ESSA tổ chức các thành phần liên "
        "lớp, Hình 4 minh họa lựa chọn chọn lọc giữa MEC và base station "
        "trong việc relay video, đặt video source, MEC units, BS và user "
        "accesses trong một pipeline ra quyết định thống nhất."
    ),
    "Bảng 3": (
        "Bảng 3 dưới đây tóm lược ba công trình cross-layer tiêu biểu trong "
        "nhánh này và cho thấy rõ rằng các framework này, dù tinh xảo về "
        "mặt kỹ thuật, vẫn chủ yếu phục vụ human-centric QoE chứ chưa hướng "
        "đến machine-centric objective."
    ),
    "Bảng 4": (
        "Để hệ thống hóa các điểm khác biệt nói trên, Bảng 4 dưới đây so "
        "sánh hai bài toán dưới góc nhìn objective trên sáu khía cạnh lý "
        "thuyết quan trọng, từ khả năng bù trừ của người dùng đến độ nhạy "
        "với lỗi cục bộ và mục tiêu thiết kế codec."
    ),
    "Hình 5": (
        "Để khái quát hóa kiến trúc tổng thể được nhắc tới, Hình 5 trình "
        "bày pipeline mô phỏng/thực nghiệm cho network-aware VCM, từ dữ "
        "liệu video, ước lượng semantic importance, mã hóa/app control, "
        "mô phỏng mạng, tới inference và đánh giá QoE."
    ),
    "Hình 6": (
        "Sơ đồ kiến trúc tổng thể của framework đề xuất được trình bày ở "
        "Hình 6, trong đó sáu mô-đun chính được nối với nhau bởi feedback "
        "loop từ QoE evaluator về cả application controller lẫn network-side "
        "controller."
    ),
    "Bảng 5": (
        "Bảng 5 dưới đây tóm tắt cấu hình thực nghiệm hiện tại của pipeline "
        "CL-ROI-VCM, bao gồm dữ liệu, tác vụ máy, codec, biến mạng, action "
        "space, state space, reward và phương pháp học tăng cường."
    ),
    "Bảng 10": (
        "Để hệ thống hóa các diễn giải vừa nêu và tránh việc \"đọc quá\" kết "
        "quả thực nghiệm, Bảng 10 dưới đây trình bày một bộ kết luận học "
        "thuật đúng mức tương ứng với từng quan sát chính từ Run 3."
    ),
    "Hình A.1": (
        "Hình A.1 dưới đây minh họa cấu trúc đa thành phần của machine-centric "
        "QoE, trong đó năm yếu tố task accuracy, end-to-end delay, "
        "bitrate/goodput, reliability/loss và compute cost cùng hội tụ về một "
        "utility tổng U_machine."
    ),
    "Hình A.2": (
        "Hình A.2 trình bày so sánh trực quan giữa ba chiến lược điều khiển — "
        "app-only, net-only và joint cross-layer — cho thấy vì sao chỉ tối ưu "
        "một phía thường chưa đủ để đạt utility đầu-cuối tốt nhất."
    ),
    "Hình A.3": (
        "Hình A.3 dưới đây phác họa action-state map mà một policy RL "
        "state-adaptive lý tưởng cần đạt được, làm rõ điều kiện cần để chứng "
        "minh adaptive behavior một cách định lượng."
    ),
    "Hình A.4": (
        "Hình A.4 trình bày đường cong rate–u_task_gt thực đo từ Run 3 trên "
        "BDD100K kết hợp libx265, giúp so sánh trực tiếp cấu trúc rate–utility "
        "giữa codec proxy ban đầu và codec HEVC thật."
    ),
    "Bảng B.1": (
        "Bảng B.1 dưới đây tóm tắt mục tiêu chính và kết luận học thuật của "
        "bốn giai đoạn thực nghiệm đã thực hiện, làm rõ vai trò chẩn đoán và "
        "vai trò bằng chứng của từng run."
    ),
    "Bảng B.2": (
        "Bảng B.2 dưới đây trình bày so sánh chi tiết các bộ trọng số reward "
        "đã được thử nghiệm qua bốn run, kèm theo cơ sở định lượng cho việc "
        "lựa chọn từng tham số."
    ),
    "Bảng F.1": (
        "Để định lượng tác động của sự thay đổi lịch trình entropy, Bảng F.1 "
        "dưới đây so sánh tỉ lệ giữa entropy bonus và policy gradient giữa "
        "ba run, kèm theo giá trị β_ent thực tế tại các mốc episode quan trọng."
    ),
    "Bảng G.1": (
        "Bảng G.1 dưới đây liệt kê các tham số chính của môi trường mô phỏng "
        "và bộ mã hóa, bao gồm action space, state space, dải bitrate và "
        "cấu hình x265 đang sử dụng."
    ),
    "Bảng G.2": (
        "Tiếp theo, Bảng G.2 trình bày cấu hình A3C và mạng actor-critic, từ "
        "learning rate, số worker, discount factor đến các tham số entropy "
        "regularization và cấu trúc head."
    ),
    "Bảng G.3": (
        "Bảng G.3 tổng hợp các trọng số reward function và các hằng số "
        "normalization, đây là những giá trị có ảnh hưởng trực tiếp đến hành "
        "vi học của RL agent."
    ),
    "Bảng G.4": (
        "Cuối cùng, Bảng G.4 mô tả hạ tầng tính toán đã được sử dụng trong "
        "các run, từ phiên bản phần mềm đến thời gian sinh profile và huấn "
        "luyện trên 1000 episode."
    ),
}


def insert_paragraph_before_element(target_element, doc, text):
    """Create a new paragraph and place it as previous sibling of target."""
    new_p = doc.add_paragraph()
    # Copy default formatting (Normal style with first-line indent)
    new_p.paragraph_format.first_line_indent = Cm(1.0)
    run = new_p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")

    # Move new_p element to before target
    target_element.addprevious(new_p._element)
    return new_p


def main():
    doc = Document(SRC)
    caption_pattern = re.compile(r"^(Hình|Bảng)\s+([A-Z]?\.?\d+(?:\.\d+)?)\.\s")

    paragraphs = list(doc.paragraphs)
    placeholder_text = "[Hình minh họa"

    inserted = 0
    for i, p in enumerate(paragraphs):
        text = p.text.strip()
        m = caption_pattern.match(text)
        if not m:
            continue
        ref_str = f"{m.group(1)} {m.group(2)}"
        if ref_str not in INTROS:
            continue

        # Identify insertion target: place new paragraph BEFORE figure placeholder
        # (if it exists right before the caption) or before the caption itself.
        target_idx = i
        if i > 0 and placeholder_text in paragraphs[i - 1].text:
            target_idx = i - 1

        # Also check that we don't double-insert (skip if previous paragraph
        # already references this caption — covers cases where v12 already has
        # the reference)
        prev_text = paragraphs[target_idx - 1].text if target_idx > 0 else ""
        if ref_str in prev_text:
            continue

        target_element = paragraphs[target_idx]._element
        intro_text = INTROS[ref_str]
        insert_paragraph_before_element(target_element, doc, intro_text)
        inserted += 1
        print(f"Inserted intro for {ref_str} before paragraph index {target_idx}")

    doc.save(OUT)
    print(f"\nTotal intros inserted: {inserted}")
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
